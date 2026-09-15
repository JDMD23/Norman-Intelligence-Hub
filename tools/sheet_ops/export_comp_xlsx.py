"""Export Lease Comps to an Excel workbook built for running NER independently.

The markdown briefs read well but cannot be calculated against. This writes the same data as
one row per comp with every NER input in its own numeric cell, so an analyst can model
directly and JD can filter, sort and share it.

What the sheet is careful about:

  Blank stays blank. A comp with no TIA gets an empty cell, never a zero — a fabricated zero
  would silently inflate its NER, which is the single largest error the book can carry. The
  Data Flags column names every such gap so nobody has to notice an empty cell.

  Derived columns are formulas, not baked numbers, so the sheet recalculates when JD edits an
  input. Only Excel-2007-era functions are used so the file opens and computes anywhere.

  Hub NER is carried alongside an empty Analyst NER column and a variance formula. The point
  of the exercise is the comparison: where the analyst and the workbook disagree, one of them
  is wrong, and the variance column says so on entry.

Usage:  python3 tools/sheet_ops/export_comp_xlsx.py [outfile.xlsx]
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import session, get_values, headers
from export_comp_briefs import floors

FONT = 'Arial'
GREEN = '003F2D'        # CBRE green — matches the workbook and the Dashboard
INPUT_FILL = PatternFill('solid', fgColor='FFFF99')   # cells the analyst fills in
FLAG_FILL = PatternFill('solid', fgColor='FFF2CC')

# (header, width, number format, source key) — source None means a formula, set per row below
COLS = [
    ('Comp ID',                 10, None,        'Comp ID'),
    ('Date Signed',             15, None,        'Date Signed'),
    ('Tenant',                  22, None,        'Tenant'),
    ('Address',                 26, None,        'Address'),
    ('Submarket',               24, None,        'Submarket'),
    ('Floor(s)',                34, None,        '__floors'),
    ('Deal Type',               18, None,        'Deal Type'),
    ('Condition',               14, None,        'Condition'),
    ('Delivery Condition',      18, None,        'Delivery Condition'),
    ('RSF',                     11, '#,##0',     'RSF'),
    ('Seats',                    8, '#,##0',     'Seats'),
    ('Term (Years)',            12, '0.00',      'Term (Years)'),
    ('Term (Months)',           13, '0',         '__formula_months'),
    ('Rent Yrs 1-5 ($/SF)',     18, '$#,##0.00', 'Rent P1 ($/RSF, mo 1-60)'),
    ('Rent Yrs 6-10 ($/SF)',    18, '$#,##0.00', 'Rent P2 ($/RSF, mo 61-120)'),
    ('Rent Yrs 11+ ($/SF)',     18, '$#,##0.00', 'Rent P3 ($/RSF, mo 121+)'),
    ('Free Rent (Months)',      17, '0.0',       'Free Rent (months)'),
    ('TIA ($/SF)',              12, '$#,##0',    'TI $/SF'),
    ('Year 1 Gross Rent ($)',   19, '$#,##0',    '__formula_yr1'),
    ('Free Rent Value ($)',     18, '$#,##0',    '__formula_free'),
    ('TIA Total ($)',           15, '$#,##0',    '__formula_ti'),
    ('Hub NER @ 6% ($/SF)',     18, '$#,##0.00', 'NER Annuity ($/RSF/Yr) @ 6%'),
    ('Analyst NER ($/SF)',      18, '$#,##0.00', '__blank'),
    ('Variance vs Hub ($)',     18, '$#,##0.00', '__formula_var'),
    ('Data Flags',              30, None,        '__flags'),
    ('Notes',                   60, None,        'Notes'),
]

README = [
    ('Lease comp export — Norman Intelligence Hub', True),
    ('', False),
    ('One row per signed lease, exported straight from the Lease Comps tab of the workbook. '
     'Nothing here is retyped or inferred.', False),
    ('', False),
    ('HOW TO USE THIS SHEET', True),
    ('Fill in the yellow "Analyst NER ($/SF)" column. Everything else is either source data or '
     'a formula.', False),
    ('"Variance vs Hub ($)" fills itself in as soon as you enter an analyst figure. A variance '
     'of more than a dollar or so means the two models disagree about something real — check '
     'the term and the TIA first, they are where it usually is.', False),
    ('', False),
    ('THE ONE THING THAT WILL BREAK YOUR NUMBERS', True),
    ('An empty cell means nobody knows the value. It does not mean zero.', False),
    ('Two comps have no TIA on file and one of those also has no free rent figure. Entering 0 '
     'for them produces an NER that looks fine and is badly overstated. The "Data Flags" '
     'column names every comp affected — filter on it before you model.', False),
    ('', False),
    ('HOW THE HUB COMPUTES NER', True),
    ('6% annual discount rate, applied monthly, beginning of month, over the full term. Free '
     'rent and TIA are both taken at face value as an upfront cost. Leasing commissions are '
     'excluded. The result is levelized into a flat annual $/SF.', False),
    ('Rent is modelled as three flat tranches: months 1-60, 61-120, and 121 onward. A blank '
     'tranche carries the previous rate forward — so a lease with a rate in "Yrs 1-5" and '
     'nothing after it is flat for its whole term, not free after year five.', False),
    ('', False),
    ('READING THE OTHER COLUMNS', True),
    ('Deal Type — subleases, renewals and expansions are not directly comparable to new '
     'leases. 27 of the 113 rows are one of those.', False),
    ('Term (Months) is the NER input, and is the full lease term including any free months.', False),
    ('Floor(s) — "Entire" and "Partial" describe how much of the floor the tenant took.', False),
]


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else 'Lease_Comps_NER.xlsx'
    s = session()
    H = headers(s, 'Lease Comps')
    raw = get_values(s, "'Lease Comps'!A2:AS1206", render='UNFORMATTED_VALUE')
    fmt = get_values(s, "'Lease Comps'!A2:AS1206", render='FORMATTED_VALUE')

    def cell(r, hdr):
        i = 0
        for ch in H[hdr]:
            i = i * 26 + ord(ch) - 64
        return r[i-1] if i-1 < len(r) else ''

    def num(r, hdr):
        v = cell(r, hdr)
        if isinstance(v, (int, float)):
            return float(v)
        v = str(v).strip()
        if not v:
            return None
        try:
            return float(v.replace('$', '').replace(',', ''))
        except ValueError:
            return None

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lease Comps'

    thin = Side(style='thin', color='D9D9D9')
    for i, (name, width, _, _) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=GREEN)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    row = 2
    flagged = 0
    for r, fr in zip(raw, fmt):
        cid = str(cell(r, 'Comp ID')).strip()
        if not cid:
            continue
        flags = []
        if num(r, 'TI $/SF') is None:
            flags.append('TIA UNKNOWN — do not enter 0')
        if num(r, 'Free Rent (months)') is None:
            flags.append('FREE RENT UNKNOWN')
        if num(r, 'Seats') is None:
            flags.append('seats unknown')
        if flags:
            flagged += 1

        for i, (name, _, numfmt, key) in enumerate(COLS, 1):
            L = get_column_letter(i)
            v = None
            if key == '__floors':
                v = floors(cell(r, 'Floor(s)'))
            elif key == '__flags':
                v = '; '.join(flags) or None
            elif key == '__blank':
                v = None
            elif key == '__formula_months':
                v = f'=IF(L{row}="","",ROUND(L{row}*12,0))'
            elif key == '__formula_yr1':
                v = f'=IF(OR(J{row}="",N{row}=""),"",J{row}*N{row})'
            elif key == '__formula_free':
                v = f'=IF(OR(J{row}="",N{row}="",Q{row}=""),"",J{row}*N{row}*Q{row}/12)'
            elif key == '__formula_ti':
                v = f'=IF(OR(J{row}="",R{row}=""),"",J{row}*R{row})'
            elif key == '__formula_var':
                v = f'=IF(OR(W{row}="",V{row}=""),"",W{row}-V{row})'
            elif key == 'Date Signed':
                v = str(cell(fr, 'Date Signed')).strip() or None
            elif numfmt:
                v = num(r, key)                      # blank stays blank, never 0
            else:
                v = str(cell(r, key)).strip() or None

            c = ws.cell(row=row, column=i, value=v)
            c.font = Font(name=FONT, size=10)
            c.border = Border(bottom=thin)
            if numfmt:
                c.number_format = numfmt
            if key == '__blank':
                c.fill = INPUT_FILL
            if key == '__flags' and flags:
                c.fill = FLAG_FILL
                c.font = Font(name=FONT, size=10, bold=True, color='9C5700')
            if key in ('Notes', '__floors'):
                c.alignment = Alignment(wrap_text=False)
        row += 1

    last = row - 1
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{last}'

    rm = wb.create_sheet('Read Me', 0)
    rm.column_dimensions['A'].width = 110
    for i, (text, bold) in enumerate(README, 1):
        c = rm.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, size=13 if i == 1 else 10, bold=bold,
                      color=GREEN if bold else '000000')
        c.alignment = Alignment(wrap_text=True, vertical='top')
    n = len(README)
    rm.cell(row=n + 2, column=1, value='EXAMPLE — what a filled-in row looks like').font = \
        Font(name=FONT, size=10, bold=True, color=GREEN)
    rm.cell(row=n + 3, column=1,
            value='LC-0104 Suno AI · 87,176 RSF · 135-month term · $92.00 then $100.00 · '
                  '15 months free · $150 TIA · Hub NER $63.49 → enter your figure in column W, '
                  'and column X shows the gap.').font = Font(name=FONT, size=10)
    rm.cell(row=n + 5, column=1,
            value=f'{last - 1} comps exported. {flagged} carry a data flag.').font = \
        Font(name=FONT, size=10, italic=True)

    wb.save(out)
    print(f'{last - 1} comps -> {out}  ({flagged} flagged)')


if __name__ == '__main__':
    main()
