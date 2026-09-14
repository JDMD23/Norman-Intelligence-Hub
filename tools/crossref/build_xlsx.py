import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import re as _re
def sc(v):
    """strip control characters openpyxl refuses (deck cells carry stray \\x0b line breaks)"""
    return _re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', ' ', v).strip() if isinstance(v,str) else v
SC='/tmp/claude-0/-home-user-Norman-Intelligence-Hub/2568e279-3cbb-5941-be03-4e24cdf0c2fe/scratchpad'
r=json.load(open(f'{SC}/report.json'))
FONT='Arial'; GREEN='003F2D'
ECON={'RSF','Term','Rent','Free rent','TI'}
AMBER=PatternFill('solid',fgColor='FFF2CC'); RED=PatternFill('solid',fgColor='FCE4E4')

# reclassify: economics vs cosmetic only
exact=list(r['exact']); cosmetic=[]; econ=[]
for x in r['differ']:
    (econ if any(f[0] in ECON for f in x['diffs']) else cosmetic).append(x)

wb=Workbook(); ws=wb.active; ws.title='Summary'
S=[('Deck cross-reference — 6 PowerPoints vs the Intelligence Hub',True),('',False),
   (f"304 rows across 6 decks -> {len(exact)+len(cosmetic)+len(econ)+len(r['missing'])} distinct transactions.",False),('',False),
   ('THE THREE BUCKETS',True),
   (f"Exact — every field agrees: {len(exact)}",False),
   (f"Cosmetic only — economics agree, a name/floor/label differs: {len(cosmetic)}",False),
   (f"Economics differ — RSF, term, rent, free rent or TI disagrees: {len(econ)}",False),
   (f"Not in the hub at all: {len(r['missing'])}",False),('',False),
   ('HOW A MATCH WAS MADE',True),
   ('Tenant name similarity is required, plus either the same building or an RSF within 5%. '
    'Building + RSF + date alone is NOT enough: Veeva Systems and Altana AI took near-identical '
    'floor plates at 2 Penn Plaza in the same month, and an earlier version of this analysis '
    'wrongly merged them.',False),('',False),
   ('WHAT THIS SETTLES',True),
   ('The open NER question was whether a "10-year lease with 16 months free" is a 120-month term '
    'or a 136-month term. Of the 45 matched comps, 28 have a deck term equal to the hub term '
    '(free rent sits INSIDE), and 9 equal hub term plus free rent (free rent sits ON TOP). '
    'The decks favour the hub roughly 3 to 1, so the analyst\'s blanket "add free rent to term" '
    'is not supported. The 9 exceptions are listed on the Economics Differ tab and look like '
    'hub entry errors rather than a convention.',False),('',False),
   ('NOTHING HAS BEEN CHANGED IN THE WORKBOOK.',True),
   ('Every row below is a proposal. Decide per comp, then tell me which to apply.',False)]
ws.column_dimensions['A'].width=118
for i,(t,b) in enumerate(S,1):
    c=ws.cell(row=i,column=1,value=sc(t))
    c.font=Font(name=FONT,size=13 if i==1 else 10,bold=b,color=GREEN if b else '000000')
    c.alignment=Alignment(wrap_text=True,vertical='top')

def sheet(name,cols,widths):
    w=wb.create_sheet(name)
    for i,(h,wd) in enumerate(zip(cols,widths),1):
        c=w.cell(row=1,column=i,value=h)
        c.font=Font(name=FONT,size=10,bold=True,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor=GREEN)
        c.alignment=Alignment(vertical='center',wrap_text=True)
        w.column_dimensions[get_column_letter(i)].width=wd
    w.row_dimensions[1].height=28
    return w

w=sheet('Exact',['Comp ID','Tenant','Address','RSF','Hub date','Deck date','Decks'],
        [10,24,28,12,15,12,44])
for i,x in enumerate(exact,2):
    for j,v in enumerate([x['cid'],x['tenant'],x['address'],x['hub_rsf'],x['date'],x['deck_date'],
                          '; '.join(x['decks'])],1):
        c=w.cell(row=i,column=j,value=sc(v)); c.font=Font(name=FONT,size=10)
        if j==4: c.number_format='#,##0'
w.auto_filter.ref=f'A1:G{len(exact)+1}'; w.freeze_panes='C2'

w=sheet('Economics Differ',['Comp ID','Tenant','Address','Field','Hub says','Decks say',
                            'Decks disagree with each other','Decks'],[10,22,24,12,30,42,40,34])
row=2
for x in econ:
    for f,hub,deck in x['diffs']:
        if f not in ECON: continue
        for j,v in enumerate([x['cid'],x['tenant'],x['address'],f,hub,deck,
                              '; '.join(x['conflicts']),'; '.join(x['decks'])],1):
            c=w.cell(row=row,column=j,value=sc(v)); c.font=Font(name=FONT,size=10)
            c.alignment=Alignment(wrap_text=True,vertical='top')
            if j in (5,6): c.fill=AMBER
            if j==7 and x['conflicts']: c.fill=RED
        row+=1
w.auto_filter.ref=f'A1:H{row-1}'; w.freeze_panes='D2'

w=sheet('Cosmetic Only',['Comp ID','Tenant','Address','Field','Hub says','Decks say'],
        [10,22,26,14,30,42])
row=2
for x in cosmetic:
    for f,hub,deck in x['diffs']:
        for j,v in enumerate([x['cid'],x['tenant'],x['address'],f,hub,deck],1):
            c=w.cell(row=row,column=j,value=sc(v)); c.font=Font(name=FONT,size=10)
            c.alignment=Alignment(wrap_text=True,vertical='top')
        row+=1
w.auto_filter.ref=f'A1:F{row-1}'; w.freeze_panes='D2'

mis=sorted([g[0] for g in r['missing']],key=lambda x:-(x['rsf'] or 0))
w=sheet('Not In Hub',['Tenant','Address','Floor','RSF','Signed','Term','Base Rent','Free Rent',
                      'TI','Deal Type','Deck','Add?'],[26,28,14,11,10,10,24,12,22,16,34,8])
for i,m in enumerate(mis,2):
    for j,v in enumerate([m['tenant'],m['address'],m['floor'],m['rsf'],m['date_raw'],m['term_raw'],
                          m['rent_raw'],m['free_raw'],m['ti_raw'],m['deal'],m['deck'],''],1):
        c=w.cell(row=i,column=j,value=sc(v)); c.font=Font(name=FONT,size=10)
        if j==4: c.number_format='#,##0'
        if j==12: c.fill=PatternFill('solid',fgColor='FFFF99')
w.auto_filter.ref=f'A1:L{len(mis)+1}'; w.freeze_panes='C2'

out=f'{SC}/Deck_Crossref.xlsx'
wb.save(out)
print(f'exact {len(exact)} | cosmetic {len(cosmetic)} | economics differ {len(econ)} | missing {len(mis)}')
print(f'saved {out}')
